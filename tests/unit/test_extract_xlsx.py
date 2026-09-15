import zipfile

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

from pii_redact.extract.xlsx import XlsxExtractionError, XlsxExtractor
from pii_redact.types import DocFormat


def _blocks_by_ref(extracted):
    return {b.source_ref: b.text for b in extracted.blocks}


def _patch_formula_cached_value(path, sheet_file, old_cell_xml, new_cell_xml):
    """openpyxl never writes a cached value for a formula it creates itself
    (confirmed by direct probing), so the only way to produce a realistic
    fixture - a workbook Excel actually opened and saved, which DOES cache
    formula results - is to patch the raw sheet XML after the fact."""
    with zipfile.ZipFile(path) as zin:
        contents = {n: zin.read(n) for n in zin.namelist()}
    xml = contents[sheet_file].decode("utf-8")
    assert old_cell_xml in xml, f"fixture assumption broken: {old_cell_xml!r} not found in {xml!r}"
    contents[sheet_file] = xml.replace(old_cell_xml, new_cell_xml).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in contents.items():
            zout.writestr(name, data)


def test_literal_string_cells_are_extracted(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["B1"] = "ABCDE1234F"
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    assert extracted.doc_format == DocFormat.XLSX
    refs = _blocks_by_ref(extracted)
    assert refs[("value", "Sheet", "A1")] == "Rahul Kumar"
    assert refs[("value", "Sheet", "B1")] == "ABCDE1234F"


def test_numeric_cells_are_skipped(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 123456789012
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    assert extracted.blocks == []


def test_empty_cells_are_skipped(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["B1"] = None
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    assert len(extracted.blocks) == 1


def test_hidden_sheet_and_hidden_row_are_included(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1.row_dimensions[3].hidden = True
    ws1["A3"] = "HiddenRowValue"

    ws2 = wb.create_sheet("Secret")
    ws2.sheet_state = "hidden"
    ws2["A1"] = "SecretSheetValue"

    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    texts = {b.text for b in extracted.blocks}
    assert "HiddenRowValue" in texts
    assert "SecretSheetValue" in texts


def test_formula_text_itself_is_never_extracted(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["A2"] = "=A1"
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    # No block's text should ever be the raw formula string.
    assert all(not b.text.startswith("=") for b in extracted.blocks)
    refs = _blocks_by_ref(extracted)
    assert ("value", "Sheet", "A2") not in refs


def test_formula_without_cached_value_yields_no_block(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["A2"] = "=A1"
    wb.save(p)  # openpyxl never caches formula results - documented gap

    extracted = XlsxExtractor().extract(p)
    refs = _blocks_by_ref(extracted)
    assert ("formula_result", "Sheet", "A2") not in refs


def test_formula_with_cached_value_is_extracted_as_read_only(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["A2"] = "=A1"
    wb.save(p)

    _patch_formula_cached_value(
        p,
        "xl/worksheets/sheet1.xml",
        '<c r="A2"><f>A1</f><v /></c>',
        '<c r="A2" t="str"><f>A1</f><v>Rahul Kumar</v></c>',
    )

    extracted = XlsxExtractor().extract(p)
    formula_blocks = [b for b in extracted.blocks if b.source_ref == ("formula_result", "Sheet", "A2")]
    assert len(formula_blocks) == 1
    assert formula_blocks[0].text == "Rahul Kumar"
    assert formula_blocks[0].read_only is True


def test_cell_comment_is_extracted(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "unrelated"
    ws["A1"].comment = Comment("Contains PAN ABCDE1234F", "author")
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    refs = _blocks_by_ref(extracted)
    assert refs[("comment", "Sheet", "A1")] == "Contains PAN ABCDE1234F"


def test_workbook_defined_name_is_extracted(tmp_path):
    p = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    wb.defined_names["Rahul_Kumar_Salary"] = DefinedName(
        "Rahul_Kumar_Salary", attr_text="Sheet!$A$1"
    )
    wb.save(p)

    extracted = XlsxExtractor().extract(p)
    refs = _blocks_by_ref(extracted)
    assert refs[("defined_name", None, "Rahul_Kumar_Salary")] == "Rahul_Kumar_Salary"
    block = next(b for b in extracted.blocks if b.source_ref == ("defined_name", None, "Rahul_Kumar_Salary"))
    assert block.read_only is True


def test_invalid_file_raises(tmp_path):
    p = tmp_path / "not_really.xlsx"
    p.write_bytes(b"this is not a zip file at all")
    with pytest.raises(XlsxExtractionError):
        XlsxExtractor().extract(p)
