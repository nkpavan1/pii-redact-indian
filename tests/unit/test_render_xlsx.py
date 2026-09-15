import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

from pii_redact.extract.xlsx import XlsxExtractor
from pii_redact.render.xlsx import XlsxRenderError, XlsxRenderer


def _block_index(extracted, source_ref):
    return next(i for i, b in enumerate(extracted.blocks) if b.source_ref == source_ref)


def test_value_round_trip_replaces_only_targeted_cell(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["B1"] = "ABCDEF1234F"
    wb.save(src)

    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("value", "Sheet", "B1"))

    out = tmp_path / "out.xlsx"
    XlsxRenderer().render(src, extracted, {idx: "IN_PAN_A"}, out)

    result = openpyxl.load_workbook(out)["Sheet"]
    assert result["A1"].value == "Rahul Kumar"
    assert result["B1"].value == "IN_PAN_A"


def test_no_replacements_leaves_cells_unchanged(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    wb.save(src)
    extracted = XlsxExtractor().extract(src)

    out = tmp_path / "out.xlsx"
    XlsxRenderer().render(src, extracted, {}, out)

    assert openpyxl.load_workbook(out)["Sheet"]["A1"].value == "Rahul Kumar"


def test_comment_round_trip(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "unrelated"
    ws["A1"].comment = Comment("Contains PAN ABCDE1234F", "author")
    wb.save(src)

    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("comment", "Sheet", "A1"))

    out = tmp_path / "out.xlsx"
    XlsxRenderer().render(src, extracted, {idx: "REDACTED"}, out)

    result = openpyxl.load_workbook(out)["Sheet"]
    assert result["A1"].value == "unrelated"  # cell value untouched
    assert result["A1"].comment.text == "REDACTED"


def test_formula_result_block_refuses_to_be_overwritten(tmp_path):
    import zipfile

    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    ws["A2"] = "=A1"
    wb.save(src)

    with zipfile.ZipFile(src) as zin:
        contents = {n: zin.read(n) for n in zin.namelist()}
    xml = contents["xl/worksheets/sheet1.xml"].decode("utf-8")
    xml = xml.replace(
        '<c r="A2"><f>A1</f><v /></c>',
        '<c r="A2" t="str"><f>A1</f><v>Rahul Kumar</v></c>',
    )
    contents["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in contents.items():
            zout.writestr(name, data)

    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("formula_result", "Sheet", "A2"))

    out = tmp_path / "out.xlsx"
    with pytest.raises(XlsxRenderError):
        XlsxRenderer().render(src, extracted, {idx: "REDACTED"}, out)


def test_defined_name_block_refuses_to_be_overwritten(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    wb.defined_names["Rahul_Kumar_Salary"] = DefinedName(
        "Rahul_Kumar_Salary", attr_text="Sheet!$A$1"
    )
    wb.save(src)

    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("defined_name", None, "Rahul_Kumar_Salary"))

    out = tmp_path / "out.xlsx"
    with pytest.raises(XlsxRenderError):
        XlsxRenderer().render(src, extracted, {idx: "PERSON_A"}, out)


def test_metadata_is_scrubbed_even_with_no_replacements(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    wb.properties.creator = "Rahul Kumar"
    wb.properties.lastModifiedBy = "Priya Singh"
    wb.properties.title = "Personal Tax Statement"
    wb.save(src)

    extracted = XlsxExtractor().extract(src)
    out = tmp_path / "out.xlsx"
    XlsxRenderer().render(src, extracted, {}, out)

    props = openpyxl.load_workbook(out).properties
    # openpyxl always re-defaults `creator` to the literal "openpyxl" on
    # save if cleared (verified quirk - see render/xlsx.py docstring); the
    # original PII value being gone is what matters, not a literal None.
    assert props.creator == "openpyxl"
    assert props.lastModifiedBy is None
    assert props.title is None


def test_hidden_sheet_and_row_state_survives_render(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1.row_dimensions[3].hidden = True
    ws1["A3"] = "HiddenRowValue"
    ws2 = wb.create_sheet("Secret")
    ws2.sheet_state = "hidden"
    ws2["A1"] = "SecretSheetValue"
    wb.save(src)

    extracted = XlsxExtractor().extract(src)
    out = tmp_path / "out.xlsx"
    XlsxRenderer().render(src, extracted, {}, out)

    result = openpyxl.load_workbook(out)
    assert result["Secret"].sheet_state == "hidden"
    assert result["Visible"].row_dimensions[3].hidden is True


def test_stale_value_cell_raises_instead_of_silently_writing_wrong_cell(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rahul Kumar"
    wb.save(src)
    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("value", "Sheet", "A1"))

    # Simulate the source changing between extract and render.
    wb2 = openpyxl.load_workbook(src)
    wb2["Sheet"]["A1"] = 12345
    wb2.save(src)

    out = tmp_path / "out.xlsx"
    with pytest.raises(XlsxRenderError):
        XlsxRenderer().render(src, extracted, {idx: "X"}, out)


def test_stale_sheet_raises(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Original"
    ws["A1"] = "Rahul Kumar"
    wb.save(src)
    extracted = XlsxExtractor().extract(src)
    idx = _block_index(extracted, ("value", "Original", "A1"))

    wb2 = openpyxl.load_workbook(src)
    wb2["Original"].title = "Renamed"
    wb2.save(src)

    out = tmp_path / "out.xlsx"
    with pytest.raises(XlsxRenderError):
        XlsxRenderer().render(src, extracted, {idx: "X"}, out)


def test_invalid_file_on_reread_raises(tmp_path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Rahul Kumar"
    wb.save(src)
    extracted = XlsxExtractor().extract(src)

    src.write_bytes(b"not a valid xlsx anymore")

    out = tmp_path / "out.xlsx"
    with pytest.raises(XlsxRenderError):
        XlsxRenderer().render(src, extracted, {}, out)
