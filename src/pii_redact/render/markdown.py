"""Markdown output for every input format - what LLMs and the wiki consume.

Uses the same `replacements` the native renderers use, so the codes are
identical whichever output is chosen. Per format:

- Text/markdown: the source with replacements spliced back in place
  (render/text.py), so frontmatter, headings and code blocks survive as-is.
- PDF and images: the extracted lines in reading order, a `## Page N`
  heading per PDF page. Lines that sit side by side on one row are joined
  with " | " so label/value pairs stay together. No PDF font fitting
  applies here, so a long code never degrades to a generic "[REDACTED]"
  marker the way it can in a redacted PDF (see render/pdf.py).
- CSV: a markdown table (the header row if the file has one, otherwise
  spreadsheet-style column letters).
- XLSX: one table per sheet with every non-empty row, cell comments listed
  under it. Formula cells show their cached result - and, unlike the native
  XLSX output, a formula result that contains PII IS replaced here, since
  nothing about the workbook's calculations depends on this text.
- JSON: the redacted document, pretty-printed in a fenced code block.

Deliberately NOT rendered, because nothing in them is scanned for PII:
sheet names (sheets are headed `## Sheet N`), defined names, PDF metadata
and annotations, and anything that is an image rather than text inside a
PDF. The file name is never written into the output either.

Known gaps carried over from extraction: numeric cells/values are not
scanned (a numeric account number is shown as-is), nor are JSON object
keys.
"""

from __future__ import annotations

import datetime as dt
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from pii_redact.render.base import Renderer
from pii_redact.render.csv_ import redacted_rows
from pii_redact.render.json_ import redacted_document
from pii_redact.render.text import redacted_text
from pii_redact.types import DocFormat, ExtractedDocument

_ROW_OVERLAP = 0.5


class MarkdownRenderError(ValueError):
    pass


def _plain(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _escape(text: str) -> str:
    """Keeps a value inside one markdown table cell / one list line."""
    return text.replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")


def _table(header: list[str], rows: list[list[str]]) -> str:
    width = max([len(header), *(len(r) for r in rows)])
    header = header + [get_column_letter(i + 1) for i in range(len(header), width)]
    lines = [
        "| " + " | ".join(_escape(h) for h in header) + " |",
        "|" + "---|" * width,
    ]
    for row in rows:
        padded = list(row) + [""] * (width - len(row))
        lines.append("| " + " | ".join(_escape(v) for v in padded) + " |")
    return "\n".join(lines)


def _fence(content: str, info: str = "") -> str:
    longest = max((len(run) for run in re.findall(r"`+", content)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}{info}\n{content}\n{fence}"


def _positioned(extracted: ExtractedDocument, replacements: dict[int, str]) -> str:
    """Lines in reading order: top to bottom, and left to right within a row
    (blocks whose vertical extents overlap enough count as one row)."""
    pages: dict[int | None, list[tuple[tuple | None, int, str]]] = {}
    for index, block in enumerate(extracted.blocks):
        text = " ".join(replacements.get(index, block.text).split())
        if text:
            pages.setdefault(block.location.page, []).append((block.location.bbox, index, text))

    sections = []
    for page in sorted(pages, key=lambda p: -1 if p is None else p):
        positioned = sorted(
            (item for item in pages[page] if item[0] is not None), key=lambda item: (item[0][1], item[0][0])
        )
        rows: list[list[tuple]] = []
        for item in positioned:
            y0, y1 = item[0][1], item[0][3]
            if rows:
                row_y0 = min(i[0][1] for i in rows[-1])
                row_y1 = max(i[0][3] for i in rows[-1])
                overlap = min(y1, row_y1) - max(y0, row_y0)
                smaller = min(y1 - y0, row_y1 - row_y0)
                if smaller > 0 and overlap / smaller >= _ROW_OVERLAP:
                    rows[-1].append(item)
                    continue
            rows.append([item])
        lines = [" | ".join(i[2] for i in sorted(row, key=lambda i: i[0][0])) for row in rows]
        # Blocks without a position (not expected from today's extractors)
        # keep their extraction order, after the positioned ones.
        lines += [text for bbox, _, text in pages[page] if bbox is None]
        body = "\n".join(lines)
        sections.append(body if page is None else f"## Page {page + 1}\n\n{body}")
    return "\n\n".join(sections)


def _xlsx(source_path: Path, extracted: ExtractedDocument, replacements: dict[int, str]) -> str:
    cell_text: dict[tuple[str, str], str] = {}
    comment_text: dict[tuple[str, str], str] = {}
    for index, new_text in replacements.items():
        kind, sheet, ref = extracted.blocks[index].source_ref
        if kind in ("value", "formula_result"):
            cell_text[(sheet, ref)] = new_text
        elif kind == "comment":
            comment_text[(sheet, ref)] = new_text
        # defined names are never rendered, so there is nothing to replace

    try:
        wb = openpyxl.load_workbook(source_path, data_only=True)
    except Exception as exc:
        raise MarkdownRenderError(f"{source_path}: could not open workbook ({exc})") from exc

    sections = []
    for number, ws in enumerate(wb.worksheets, start=1):
        rows, comments, last_col, letters = [], [], 0, []
        for row in ws.iter_rows():
            values = []
            for cell in row:
                key = (ws.title, cell.coordinate)
                values.append(cell_text.get(key, _plain(cell.value)))
                comment = getattr(cell, "comment", None)  # merged cells have none
                if comment is not None and comment.text:
                    comments.append(f"- {cell.coordinate}: " + _escape(comment_text.get(key, comment.text)))
            if any(v != "" for v in values):
                rows.append(values)
                last_col = max(last_col, max(i for i, v in enumerate(values) if v != "") + 1)
                letters = [get_column_letter(c.column) for c in row]
        if rows:
            body = _table(letters[:last_col], [r[:last_col] for r in rows])
        else:
            body = "_(empty sheet)_"
        if comments:
            body += "\n\nComments:\n\n" + "\n".join(comments)
        sections.append(f"## Sheet {number}\n\n{body}")
    return "\n\n".join(sections)


def to_markdown(source_path: Path, extracted: ExtractedDocument, replacements: dict[int, str]) -> str:
    fmt = extracted.doc_format
    if fmt == DocFormat.TEXT:
        return redacted_text(source_path, extracted, replacements)
    if fmt in (DocFormat.PDF, DocFormat.IMAGE):
        body = _positioned(extracted, replacements)
    elif fmt == DocFormat.CSV:
        header, rows = redacted_rows(source_path, extracted, replacements)
        body = _table(header or [], rows) if (header or rows) else ""
    elif fmt == DocFormat.XLSX:
        body = _xlsx(source_path, extracted, replacements)
    elif fmt == DocFormat.JSON:
        document = redacted_document(source_path, extracted, replacements)
        body = _fence(json.dumps(document, indent=2, ensure_ascii=False), "json")
    else:
        raise MarkdownRenderError(f"no markdown rendering for {fmt}")
    return body + "\n"


class MarkdownRenderer(Renderer):
    def render(
        self,
        source_path: Path,
        extracted: ExtractedDocument,
        replacements: dict[int, str],
        output_path: Path,
    ) -> None:
        text = to_markdown(source_path, extracted, replacements)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(text, encoding="utf-8", newline="")
