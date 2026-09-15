"""XLSX extraction via openpyxl.

Iterates every worksheet - hidden or not. openpyxl's normal cell/row
iteration already includes hidden and grouped sheets, rows, and columns;
nothing special is required to "include" them, only discipline to never
add a visibility filter later (see test_hidden_sheet_and_row_are_included
for the regression guard).

Classifies each cell as a formula or a literal value, since a formula's own
text (e.g. "=A1&B1") is not PII and must never be scanned or overwritten as
one - see project instructions, XLSX section.

Formula cells: the formula text itself is skipped for detection. Their last
*cached* computed result is extracted separately (as a `read_only=True`
block) from a second, data_only=True load - openpyxl can't expose both the
formula text and its cached value from a single load. This means a
formula's displayed result still surfaces in the review-gate preview if it
contains PII, but is clearly marked as something the render stage must
never overwrite in place (overwriting a formula cell destroys the
calculation - see render/xlsx.py). If the workbook was written by a tool
that never computed/cached formula results (e.g. openpyxl itself, or
pandas.to_excel) there is no cached value to find - a documented gap, not a
bug: this project doesn't implement a formula-evaluation engine.

Cell comments and defined names are extracted too (project instructions:
both "can carry PII outside the visible grid"). A defined name's
identifier (e.g. someone named a range "Rahul_Kumar_Salary") is scanned;
its target reference (e.g. "Sheet1!$B$2") is not itself PII text - the
referenced cell's own value is already captured as a normal cell block.
Defined names are marked `read_only=True`, same as formula results:
renaming one safely requires validating the replacement against Excel's
identifier rules (no spaces, can't start with a digit, can't collide with
a cell reference, etc.), which render/xlsx.py doesn't implement - see its
module docstring.

Workbook document properties (author, company, ...) are NOT extracted here
- see render/xlsx.py, which owns metadata scrubbing as a separate concern.

Known, documented limitations (same pattern as extract/json_.py):
- Only string-typed cell values are extracted. A PAN or account number
  entered as a numeric cell will not be caught - treating every numeric
  cell as a PII candidate would flood detection with legitimate amounts,
  dates-as-serials, and quantities.
- Whether a literal cell is itself an input to some formula elsewhere in
  the workbook is NOT determined here - that requires parsing and
  resolving cell references across every formula in the workbook, a
  significant undertaking on its own (see project instructions: "decide
  per-column whether a cell is a formula input... don't just overwrite in
  place"). Until that exists, treat any workbook containing formulas as
  needing manual review before trusting automatic overwrite of literal
  cells.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from pii_redact.extract.base import Extractor
from pii_redact.types import DocFormat, ExtractedDocument, Location, TextBlock


class XlsxExtractionError(ValueError):
    pass


def _cached_values_by_sheet(path: Path) -> dict[str, Worksheet]:
    try:
        wb_values = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise XlsxExtractionError(
            f"{path}: could not open workbook for cached-value read ({exc})"
        ) from exc
    return {ws.title: ws for ws in wb_values.worksheets}


def _extract_cell(ws: Worksheet, cell, cached_ws: Worksheet | None, blocks: list[TextBlock]) -> None:
    if cell.data_type == "f":
        cached_value = cached_ws[cell.coordinate].value if cached_ws is not None else None
        if isinstance(cached_value, str) and cached_value.strip():
            blocks.append(
                TextBlock(
                    text=cached_value,
                    location=Location(sheet=ws.title, cell=cell.coordinate),
                    source_ref=("formula_result", ws.title, cell.coordinate),
                    read_only=True,
                )
            )
        # else: no cached value available - documented gap, see module docstring.
    elif isinstance(cell.value, str) and cell.value.strip():
        blocks.append(
            TextBlock(
                text=cell.value,
                location=Location(sheet=ws.title, cell=cell.coordinate),
                source_ref=("value", ws.title, cell.coordinate),
            )
        )

    comment = cell.comment
    if comment is not None and comment.text and comment.text.strip():
        blocks.append(
            TextBlock(
                text=comment.text,
                location=Location(sheet=ws.title, cell=cell.coordinate),
                source_ref=("comment", ws.title, cell.coordinate),
            )
        )


class XlsxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedDocument:
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
        except Exception as exc:
            raise XlsxExtractionError(f"{path}: could not open workbook ({exc})") from exc

        cached_sheets = _cached_values_by_sheet(path)

        blocks: list[TextBlock] = []
        for ws in wb.worksheets:
            cached_ws = cached_sheets.get(ws.title)
            for row in ws.iter_rows():
                for cell in row:
                    _extract_cell(ws, cell, cached_ws, blocks)

            for name in getattr(ws, "defined_names", {}):
                blocks.append(
                    TextBlock(
                        text=name,
                        location=Location(sheet=ws.title, cell=None),
                        source_ref=("defined_name", ws.title, name),
                        read_only=True,
                    )
                )

        for name in wb.defined_names:
            blocks.append(
                TextBlock(
                    text=name,
                    location=Location(sheet=None, cell=None),
                    source_ref=("defined_name", None, name),
                    read_only=True,
                )
            )

        return ExtractedDocument(
            doc_format=DocFormat.XLSX,
            source_path=path,
            blocks=blocks,
        )
